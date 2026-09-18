"""Genera la plantilla base `app/templates/informe_base.docx` usada por
docxtpl para renderizar el informe final (Etapa 4).

Se ejecuta una sola vez (o cuando cambie la estructura fija del informe) y
el resultado se versiona en el repositorio como cualquier otro archivo de
plantilla. No se ejecuta en cada generación de informe.

Uso: python scripts/generar_plantilla_base.py
"""

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.reportes import estilos  # noqa: E402
from app.reportes.colores_revision import LEYENDA, sombrear_run  # noqa: E402

SALIDA = Path(__file__).resolve().parent.parent / "app" / "templates" / "informe_base.docx"

# Marcadores de texto que generador.py localiza y reemplaza por las tablas
# de puntajes/porcentajes construidas con python-docx (más simple y robusto
# que las etiquetas de fila de docxtpl para tablas puramente numéricas).
MARCADOR_TABLA_PUNTAJES = "[[TABLA_PUNTAJES]]"
MARCADOR_TABLA_PORCENTAJES = "[[TABLA_PORCENTAJES]]"


def _seccion(doc, titulo: str, *, nivel: int = 1):
    """Encabezado de sección con la regla de color debajo (solo en nivel 1,
    para marcar el inicio de los bloques principales del informe)."""
    encabezado = doc.add_heading(titulo, level=nivel)
    if nivel == 1:
        estilos.agregar_regla_inferior(encabezado)
    return encabezado


def _parrafo_justificado(doc, *runs_con_texto):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for texto in runs_con_texto:
        p.add_run(texto)
    return p


def _caption(doc, texto: str):
    p = doc.add_paragraph(texto)
    estilos.estilizar_texto_cita(p)
    return p


def construir_documento() -> Document:
    doc = Document()
    estilos.aplicar_estilos_base(doc)
    estilos.configurar_pie_pagina(doc.sections[0])

    # --- Portada ---------------------------------------------------------
    titulo = doc.add_heading("Informe de Pre Diagnóstico Empresarial", level=0)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    estilos.agregar_regla_inferior(titulo, grosor=12)

    subtitulo = doc.add_paragraph()
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitulo.add_run(
        "Identificación de factores de riesgo ergonómico por manejo manual de cargas, "
        "con base en la NOM-036-1-STPS-2018"
    )
    r.font.size = Pt(13)
    r.font.color.rgb = estilos.COLOR_TEXTO_TENUE
    r.font.italic = True

    doc.add_paragraph()
    ficha = doc.add_paragraph()
    ficha.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ficha.add_run("Empresa evaluada: ").bold = True
    ficha.add_run("{{r empresa_nombre }}")
    ficha_fecha = doc.add_paragraph()
    ficha_fecha.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ficha_fecha.add_run("Fecha de generación: ").bold = True
    ficha_fecha.add_run("{{ fecha_generacion }}")

    aviso = doc.add_paragraph()
    aviso.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_aviso = aviso.add_run(
        "Documento confidencial. Su contenido está dirigido exclusivamente a la empresa "
        "evaluada y al equipo responsable de la validación de este instrumento."
    )
    r_aviso.font.size = Pt(9.5)
    r_aviso.font.italic = True
    r_aviso.font.color.rgb = estilos.COLOR_TEXTO_TENUE

    doc.add_page_break()

    # --- Introducción ------------------------------------------------------
    _seccion(doc, "Introducción")
    _parrafo_justificado(
        doc,
        "La NOM-036-1-STPS-2018 busca identificar y prevenir los riesgos de salud que "
        "puede causar el manejo manual de cargas en el trabajo. Este documento resume "
        "los resultados de la autoevaluación que respondió la empresa ",
        "{{r empresa_nombre }}.",
    )

    doc.add_paragraph("{% if es_revision %}")
    _seccion(doc, "Guía de colores de este documento", nivel=2)
    doc.add_paragraph(
        "Esta versión del informe, generada durante la etapa de validación piloto, resalta el "
        "texto según su origen dentro del motor de recomendaciones, para facilitar su revisión:"
    )
    for etiqueta, descripcion, color in LEYENDA:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(etiqueta)
        r.bold = True
        sombrear_run(r, color)
        p.add_run(f": {descripcion}")
    doc.add_paragraph("{% endif %}")

    # --- Información general de la empresa ---------------------------------
    _seccion(doc, "Información general de la empresa")
    p = _parrafo_justificado(doc, "La empresa ", "{{r empresa_nombre }}")
    p.add_run(" se encuentra ubicada en {{r empresa_ubicacion }}. Cuenta con aproximadamente ")
    p.add_run("{{r empresa_num_trabajadores }}")
    p.add_run(
        " trabajadores, con el siguiente esquema de turnos: {{r empresa_turnos }}. Su giro "
        "es {{r empresa_giro }}. Las actividades relacionadas con el manejo manual de "
        "cargas identificadas son: {{r empresa_descripcion_mmh }}."
    )

    # --- Resultados de la evaluación ----------------------------------------
    _seccion(doc, "Resultados de la evaluación")
    _parrafo_justificado(
        doc,
        "El cuestionario se divide en temas. Cada tema tiene diez preguntas, calificadas "
        "del 0 (Nada) al 4 (Óptimo). Entre más alto es el puntaje de un tema, mejor cumple "
        "la empresa con la norma en ese tema.",
    )

    doc.add_paragraph("Tabla 1. Calificación obtenida por criterio evaluado.")
    doc.add_paragraph(MARCADOR_TABLA_PUNTAJES)

    doc.add_paragraph("Tabla 2. Porcentaje de cumplimiento por criterio evaluado.")
    doc.add_paragraph(MARCADOR_TABLA_PORCENTAJES)

    img_p = doc.add_paragraph("{{ grafica_radar }}")
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _caption(doc, "Figura 1. Porcentaje de cumplimiento por criterio, sobre las 5 franjas de severidad.")

    # --- Resultados por criterio ---------------------------------------------
    _seccion(doc, "Resultados por criterio")
    doc.add_paragraph("{% for c in criterios %}")
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    cp.add_run("Criterio {{ c.numero }}.- ").bold = True
    cp.add_run("{{r c.narrativa }}")
    doc.add_paragraph("{% endfor %}")

    # --- Áreas de oportunidad --------------------------------------------
    _seccion(doc, "Áreas de oportunidad")
    _parrafo_justificado(doc, "{{r resultado_global_apertura }}")
    _parrafo_justificado(doc, "{{r resultado_global_cierre }}")

    doc.add_paragraph("{% if hay_temas %}")
    img_temas_p = doc.add_paragraph("{{ grafica_temas }}")
    img_temas_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _caption(doc, "Figura 2. Proporción de temas de atención obligatoria y optativa.")
    doc.add_paragraph("{% endif %}")

    _seccion(doc, "Temas obligatorios", nivel=2)
    doc.add_paragraph(
        "Derivados directamente de las exigencias de la NOM-036-1-STPS-2018 para los "
        "temas evaluados con nivel de cumplimiento bajo."
    )
    doc.add_paragraph("{% for tema in temas_obligatorios %}")
    tp = doc.add_paragraph(style="List Bullet")
    tp.add_run("{{ loop.index }}. {{r tema }}")
    doc.add_paragraph("{% endfor %}")

    _seccion(doc, "Temas optativos", nivel=2)
    doc.add_paragraph("Buenas prácticas sugeridas, adicionales a lo que exige la norma.")
    doc.add_paragraph("{% for tema in temas_optativos %}")
    top = doc.add_paragraph(style="List Bullet")
    top.add_run("{{ loop.index }}. {{r tema }}")
    doc.add_paragraph("{% endfor %}")

    cierre_regla = doc.add_paragraph()
    estilos.agregar_regla_inferior(cierre_regla, grosor=4)
    doc.add_paragraph("Atentamente,")
    firma = doc.add_paragraph("Equipo de investigación en Ergonomía y Factores Humanos")
    firma.runs[0].bold = True

    return doc


if __name__ == "__main__":
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    construir_documento().save(SALIDA)
    print(f"Plantilla generada en {SALIDA}")
