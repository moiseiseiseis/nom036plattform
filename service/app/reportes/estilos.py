"""Estilos visuales del informe .docx "cuasi-final" (Etapa 6): tipografía,
color de marca, tablas y numeración de página.

Independiente de `colores_revision.py`, que solo resalta el origen del texto
durante la validación piloto y se superpone a estos estilos cuando
`modo_revision=True` — este módulo define cómo se ve el documento en
cualquier modo, incluido el informe final que se entrega a una empresa real.
"""

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from app.engine.scoring import ETIQUETA_POR_BUCKET

FUENTE = "Calibri"

COLOR_ACENTO_HEX = "1F3A5F"
COLOR_ACENTO = RGBColor(0x1F, 0x3A, 0x5F)
COLOR_TEXTO = RGBColor(0x26, 0x26, 0x26)
COLOR_TEXTO_TENUE_HEX = "595956"
COLOR_TEXTO_TENUE = RGBColor(0x59, 0x59, 0x56)
COLOR_ZEBRA_HEX = "F2F1EE"
COLOR_BLANCO = RGBColor(0xFF, 0xFF, 0xFF)

# Mismo azul que los encabezados del documento (COLOR_ACENTO) — reutilizado
# como "color de marca" de las gráficas (`grafica.py`) para que la línea de
# datos del radar se lea como parte del mismo documento, no como un elemento
# ajeno en negro.
COLOR_MARCA_HEX = COLOR_ACENTO_HEX

# Misma paleta de severidad usada en `grafica.py`, para que la clasificación
# de la tabla 2 y las bandas del radar se lean como el mismo código de color.
# `ETIQUETA_POR_BUCKET` vive en `engine/scoring.py` (junto al resto de la
# lógica de bucket) — se reexporta aquí para que quien ya importa colores de
# `estilos.py` no tenga que ir a buscar las etiquetas a otro paquete.
COLOR_POR_BUCKET = {
    "inexistente": "B91C1C",
    "minimo": "EB6834",
    "regular": "7A6900",
    "aceptable": "1BAF7A",
    "optimo": "2A78D6",
}

# Colores de "temas obligatorios vs. optativos", compartidos con la gráfica
# de pastel correspondiente.
COLOR_OBLIGATORIO_HEX = "4A3AA7"
COLOR_OPTATIVO_HEX = "898781"


def aplicar_estilos_base(doc) -> None:
    """Tipografía y color de texto por defecto para todo el documento."""
    normal = doc.styles["Normal"]
    normal.font.name = FUENTE
    normal.font.size = Pt(11)
    normal.font.color.rgb = COLOR_TEXTO
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    titulo = doc.styles["Title"]
    titulo.font.name = FUENTE
    titulo.font.size = Pt(28)
    titulo.font.color.rgb = COLOR_ACENTO
    titulo.font.bold = True

    for nombre, tamano in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 12)):
        estilo = doc.styles[nombre]
        estilo.font.name = FUENTE
        estilo.font.size = Pt(tamano)
        estilo.font.color.rgb = COLOR_ACENTO
        estilo.font.bold = True
        estilo.font.italic = False
        estilo.paragraph_format.space_before = Pt(18)
        estilo.paragraph_format.space_after = Pt(6)

    lista = doc.styles["List Bullet"]
    lista.font.name = FUENTE
    lista.font.size = Pt(11)
    lista.font.color.rgb = COLOR_TEXTO


def agregar_regla_inferior(paragraph, color_hex: str = COLOR_ACENTO_HEX, grosor: int = 6) -> None:
    """Borde inferior de color en un párrafo — usado bajo `Heading 1` para
    marcar visualmente el inicio de cada sección."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    borde = OxmlElement("w:bottom")
    borde.set(qn("w:val"), "single")
    borde.set(qn("w:sz"), str(grosor))
    borde.set(qn("w:space"), "4")
    borde.set(qn("w:color"), color_hex)
    pBdr.append(borde)
    pPr.append(pBdr)


def estilizar_texto_cita(paragraph) -> None:
    """Estilo de pie de figura / nota al pie de tabla: centrado, gris, itálica."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in paragraph.runs:
        run.font.size = Pt(9.5)
        run.font.italic = True
        run.font.color.rgb = COLOR_TEXTO_TENUE


def sombrear_celda(cell, color_hex: str) -> None:
    """Aplica un color de fondo a una celda completa de tabla."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def estilizar_tabla(tabla, *, alineacion_columnas: list[str] | None = None) -> None:
    """Aplica el look estándar de tabla del informe: encabezado en color de
    acento con texto blanco en negritas, filas de datos alternadas (zebra) y
    alineación por columna opcional (una letra por columna: 'l'/'c'/'r')."""
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER

    encabezado = tabla.rows[0]
    for cell in encabezado.cells:
        sombrear_celda(cell, COLOR_ACENTO_HEX)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = COLOR_BLANCO

    alineacion_por_letra = {
        "l": WD_ALIGN_PARAGRAPH.LEFT,
        "c": WD_ALIGN_PARAGRAPH.CENTER,
        "r": WD_ALIGN_PARAGRAPH.RIGHT,
    }
    for i, fila in enumerate(tabla.rows[1:]):
        if i % 2 == 1:
            for cell in fila.cells:
                sombrear_celda(cell, COLOR_ZEBRA_HEX)
        if alineacion_columnas:
            for cell, letra in zip(fila.cells, alineacion_columnas):
                for p in cell.paragraphs:
                    p.alignment = alineacion_por_letra[letra]


def colorear_texto(run, color_hex: str, *, negritas: bool = True) -> None:
    run.font.color.rgb = RGBColor.from_string(color_hex)
    run.font.bold = negritas


def configurar_pie_pagina(section) -> None:
    """Numeración 'Página X de Y' centrada en el pie de página, en todas las
    secciones del documento."""
    parrafo = section.footer.paragraphs[0]
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _agregar_texto_campo(parrafo, "Página ")
    _agregar_campo(parrafo, "PAGE")
    _agregar_texto_campo(parrafo, " de ")
    _agregar_campo(parrafo, "NUMPAGES")


def _agregar_texto_campo(paragraph, texto: str) -> None:
    run = paragraph.add_run(texto)
    run.font.name = FUENTE
    run.font.size = Pt(9)
    run.font.color.rgb = COLOR_TEXTO_TENUE


def _agregar_campo(paragraph, instruccion: str) -> None:
    """Inserta un campo simple de Word (`w:fldSimple`), p. ej. `PAGE` o
    `NUMPAGES`. python-docx no tiene API de alto nivel para campos; Word
    recalcula el valor cacheado al abrir o imprimir el documento."""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instruccion)
    run_elem = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), FUENTE)
    rPr.append(rFonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    rPr.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), COLOR_TEXTO_TENUE_HEX)
    rPr.append(color)
    run_elem.append(rPr)
    texto = OxmlElement("w:t")
    texto.text = "1"
    run_elem.append(texto)
    fld.append(run_elem)
    paragraph._p.append(fld)


def fijar_ancho_columnas(tabla, anchos_mm: list[float]) -> None:
    """Fija el ancho de cada columna en milímetros (Word suele ignorar
    `cell.width` si la tabla está en autoajuste; se desactiva por tabla)."""
    tabla.autofit = False
    for fila in tabla.rows:
        for cell, ancho in zip(fila.cells, anchos_mm):
            cell.width = Mm(ancho)
