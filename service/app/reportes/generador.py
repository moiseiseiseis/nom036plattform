import io
import os
from datetime import datetime
from pathlib import Path

from docx import Document as DocxDocument
from docx.document import Document as DocxDocumentType
from docx.oxml.ns import qn
from docx.shared import Mm
from docxtpl import DocxTemplate, InlineImage, RichText

from app.db import get_connection
from app.engine.engine import evaluar_criterio, evaluar_global
from app.engine.models import ResultadoCriterio, ResultadoGlobal
from app.engine.scoring import calcular_global, calcular_porcentaje, calcular_puntaje, clasificar_bucket

from . import colores_revision as colores
from . import estilos
from .grafica import generar_grafica_radar, generar_grafica_temas
from .models import DatosEvaluacion
from .repository import fetch_datos_evaluacion

PLANTILLA_BASE = Path(__file__).resolve().parent.parent / "templates" / "informe_base.docx"

MARCADOR_TABLA_PUNTAJES = "[[TABLA_PUNTAJES]]"
MARCADOR_TABLA_PORCENTAJES = "[[TABLA_PORCENTAJES]]"

BUCKET_ETIQUETAS = estilos.ETIQUETA_POR_BUCKET

MESES_ES = (
    "enero", "febrero", "marzo", "abril", "mayo", "junio",
    "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre",
)


def _fecha_generacion_es(momento: datetime) -> str:
    return f"{momento.day} de {MESES_ES[momento.month - 1]} de {momento.year}"


# Firma del responsable de la elaboración (NOM-036, numeral 7.4 inciso f):
# nombre completo + cédula profesional son datos reales de una persona, no
# contenido que este proyecto pueda inventar — se configuran por variable de
# entorno y, mientras no estén disponibles, el informe firma como equipo
# (comportamiento actual) en vez de mostrar un campo vacío o inventado.
def _responsable() -> tuple[str, str | None]:
    nombre = os.environ.get("INFORME_RESPONSABLE_NOMBRE")
    cedula = os.environ.get("INFORME_RESPONSABLE_CEDULA")
    if not nombre:
        return "Equipo de investigación en Ergonomía y Factores Humanos", None
    return nombre, cedula


def generar_informe(evaluacion_id: str, modo_revision: bool = True) -> bytes:
    """Punto de entrada usado por la API: trae los datos de la evaluación de
    la base de datos y produce el `.docx` final.

    `modo_revision` resalta el texto por su origen (dato de la empresa,
    resultado calculado, plantilla fija, recomendación de catálogo) para
    facilitar la validación piloto de la Etapa 6. Por defecto está activo
    mientras dure esa etapa; se desactivará para el informe de entrega real
    a una empresa."""
    with get_connection() as conn:
        datos = fetch_datos_evaluacion(conn, evaluacion_id)
    return generar_informe_desde_datos(datos, modo_revision=modo_revision)


def calcular_resultados(
    datos: DatosEvaluacion,
) -> tuple[list[ResultadoCriterio], ResultadoGlobal]:
    """Corre el motor de cálculo (Etapa 2) sobre los datos de una evaluación
    ya resueltos. No toca la base de datos ni genera ningún documento —
    la reutilizan tanto el generador de `.docx` como el resumen JSON del
    panel privado (Etapa 5)."""
    if not datos.criterios:
        raise ValueError("La evaluación no tiene respuestas registradas para ningún criterio")

    resultados: list[ResultadoCriterio] = []
    for dc in datos.criterios:
        puntaje, puntaje_maximo = calcular_puntaje([r.nivel for r in dc.respuestas])
        bucket = clasificar_bucket(calcular_porcentaje(puntaje, puntaje_maximo))
        resultados.append(
            evaluar_criterio(
                criterio_id=dc.id,
                numero=dc.numero,
                nombre=dc.nombre,
                respuestas=dc.respuestas,
                recomendaciones=dc.recomendaciones,
                plantilla_apertura=dc.plantillas_apertura.get(bucket, ""),
            )
        )

    _, _, _, bucket_global = calcular_global([(r.puntaje, r.puntaje_maximo) for r in resultados])
    resultado_global = evaluar_global(
        resultados,
        plantilla_cierre=datos.plantillas_cierre_global.get(bucket_global, ""),
        total_criterios=datos.total_criterios,
    )
    return resultados, resultado_global


def generar_informe_desde_datos(datos: DatosEvaluacion, modo_revision: bool = True) -> bytes:
    """Lógica de ensamblado del informe a partir de datos ya resueltos, sin
    tocar la base de datos — es lo que permite probar el generador completo
    (motor + gráfica + plantilla + tablas) con datos sintéticos."""
    resultados, resultado_global = calcular_resultados(datos)

    tpl = DocxTemplate(str(PLANTILLA_BASE))
    grafica_imagen = InlineImage(tpl, io.BytesIO(generar_grafica_radar(resultados)), width=Mm(140))

    # El marcador `{{r ... }}` de la plantilla solo acepta objetos RichText
    # (necesita aislar el tag en su propio run de XML para insertar texto
    # enriquecido); por eso todo el texto que pasa por él se envuelve en
    # RichText incluso cuando no hay color que aplicar (modo_revision=False).
    color_empresa = colores.DATO_EMPRESA if modo_revision else None
    color_plantilla = colores.PLANTILLA_FIJA if modo_revision else None
    color_recomendacion = colores.RECOMENDACION_CATALOGO if modo_revision else None
    color_tema = colores.TEMA_OBLIGATORIO_OPTATIVO if modo_revision else None

    empresa_nombre = _rt(datos.empresa_nombre, color_empresa, bold=True)
    empresa_ubicacion = _rt(datos.empresa_ubicacion or "no especificada", color_empresa)
    empresa_giro = _rt(datos.empresa_giro or "no especificado", color_empresa)
    empresa_num_trabajadores = _rt(
        datos.empresa_num_trabajadores
        if datos.empresa_num_trabajadores is not None
        else "no especificado",
        color_empresa,
        bold=True,
    )
    empresa_turnos = _rt(datos.empresa_turnos or "no especificados", color_empresa)
    empresa_descripcion_mmh = _rt(datos.empresa_descripcion_mmh or "no especificadas", color_empresa)
    criterios_ctx = [
        {
            "numero": r.numero,
            "nombre": r.nombre,
            "narrativa": _narrativa_rt(r, color_plantilla, color_recomendacion),
        }
        for r in resultados
    ]
    resultado_global_apertura = _rt(
        datos.plantillas_apertura_global.get(resultado_global.bucket, ""), color_plantilla
    )
    # Con evaluación parcial, `resultado_global.cierre` ya es el párrafo
    # completo con el alcance acotado (`construir_cierre_parcial`) — mostrar
    # además la apertura fija ("A continuación se detallan...") sería
    # redundante y no tiene bucket global real del que hablar todavía.
    resultado_global_cierre = _rt(resultado_global.cierre, color_plantilla)
    responsable_nombre, responsable_cedula = _responsable()
    temas_obligatorios = [_rt(t, color_tema) for t in resultado_global.temas_obligatorios]
    temas_optativos = [_rt(t, color_tema) for t in resultado_global.temas_optativos]

    num_obligatorios = len(resultado_global.temas_obligatorios)
    num_optativos = len(resultado_global.temas_optativos)
    hay_temas = (num_obligatorios + num_optativos) > 0
    grafica_temas_imagen = (
        InlineImage(
            tpl, io.BytesIO(generar_grafica_temas(num_obligatorios, num_optativos)), width=Mm(100)
        )
        if hay_temas
        else None
    )

    tpl.render(
        {
            "empresa_nombre": empresa_nombre,
            "empresa_ubicacion": empresa_ubicacion,
            "empresa_giro": empresa_giro,
            "empresa_num_trabajadores": empresa_num_trabajadores,
            "empresa_turnos": empresa_turnos,
            "empresa_descripcion_mmh": empresa_descripcion_mmh,
            "grafica_radar": grafica_imagen,
            "criterios": criterios_ctx,
            "resultado_global_apertura": resultado_global_apertura,
            "resultado_global_cierre": resultado_global_cierre,
            "temas_obligatorios": temas_obligatorios,
            "temas_optativos": temas_optativos,
            "grafica_temas": grafica_temas_imagen,
            "hay_temas": hay_temas,
            "es_revision": modo_revision,
            "fecha_generacion": _fecha_generacion_es(datetime.now()),
            "resultado_global_es_parcial": resultado_global.es_parcial,
            "responsable_nombre": responsable_nombre,
            "responsable_cedula": responsable_cedula,
        }
    )

    buffer = io.BytesIO()
    tpl.save(buffer)
    buffer.seek(0)

    documento = DocxDocument(buffer)
    _normalizar_sombreado_richtext(documento)
    _reemplazar_marcador_con_tabla(
        documento, MARCADOR_TABLA_PUNTAJES, _construir_tabla_puntajes, resultados, modo_revision
    )
    _reemplazar_marcador_con_tabla(
        documento, MARCADOR_TABLA_PORCENTAJES, _construir_tabla_porcentajes, resultados, modo_revision
    )

    salida = io.BytesIO()
    documento.save(salida)
    return salida.getvalue()


def _normalizar_sombreado_richtext(documento: DocxDocumentType) -> None:
    """Corrige el `w:shd` que genera `docxtpl.RichText.add(highlight=...)`
    (usado en modo_revision): escribe `<w:shd w:fill="RRGGBB"/>` sin el
    atributo `w:val`, que el esquema OOXML exige — Word lo trata como
    inválido y no pinta ningún fondo, por lo que el resaltado de la guía de
    colores no se veía en los documentos exportados (retroalimentación de la
    validación piloto, `retroalimentacion/reporte.txt`, hallazgo 7). El
    sombreado de celda que arma `estilos.sombrear_celda` ya incluye estos
    atributos y no lo toca este parche."""
    for shd in documento.element.body.iter(qn("w:shd")):
        if shd.get(qn("w:val")) is None:
            shd.set(qn("w:val"), "clear")
        if shd.get(qn("w:color")) is None:
            shd.set(qn("w:color"), "auto")


def _rt(texto, color_hex: str | None, *, bold: bool = False) -> RichText:
    """Envuelve un valor en un RichText, resaltado con el color dado si se
    proporciona uno (o sin resaltar, si `color_hex` es `None`)."""
    rt = RichText()
    rt.add(texto, highlight=color_hex, bold=bold)
    return rt


def _narrativa_rt(
    resultado: ResultadoCriterio, color_plantilla: str | None, color_recomendacion: str | None
) -> RichText:
    """Narrativa de un criterio con la plantilla de apertura y los hallazgos
    resaltados con colores distintos, reproduciendo exactamente el mismo
    ensamblado de `narrativa.ensamblar_narrativa_criterio`."""
    rt = RichText()
    rt.add(resultado.plantilla_apertura, highlight=color_plantilla)
    for h in resultado.hallazgos:
        rt.add(f" {h.texto_recomendacion}", highlight=color_recomendacion)
    return rt


def _reemplazar_marcador_con_tabla(documento, marcador, construir_tabla, resultados, modo_revision):
    for parrafo in documento.paragraphs:
        if parrafo.text.strip() == marcador:
            tabla = construir_tabla(documento, resultados, modo_revision)
            parrafo._p.addnext(tabla._tbl)
            parrafo._p.getparent().remove(parrafo._p)
            return
    raise ValueError(f"No se encontró el marcador {marcador!r} en la plantilla")


def _construir_tabla_puntajes(
    documento: DocxDocumentType, resultados: list[ResultadoCriterio], modo_revision: bool
):
    tabla = documento.add_table(rows=1, cols=2)
    tabla.style = "Table Grid"
    encabezado = tabla.rows[0].cells
    encabezado[0].text = "Criterio evaluado"
    encabezado[1].text = "Puntaje obtenido"
    for r in resultados:
        fila = tabla.add_row().cells
        fila[0].text = f"{r.numero}. {r.nombre}"
        fila[1].text = f"{r.puntaje}/{r.puntaje_maximo}"
        if modo_revision:
            colores.sombrear_run(fila[1].paragraphs[0].runs[0], colores.RESULTADO_CALCULADO)
    estilos.estilizar_tabla(tabla, alineacion_columnas=["l", "c"])
    estilos.fijar_ancho_columnas(tabla, [110, 50])
    return tabla


def _construir_tabla_porcentajes(
    documento: DocxDocumentType, resultados: list[ResultadoCriterio], modo_revision: bool
):
    tabla = documento.add_table(rows=1, cols=3)
    tabla.style = "Table Grid"
    encabezado = tabla.rows[0].cells
    encabezado[0].text = "Criterio evaluado"
    encabezado[1].text = "Porcentaje de cumplimiento"
    encabezado[2].text = "Clasificación"
    for r in resultados:
        fila = tabla.add_row().cells
        fila[0].text = f"{r.numero}. {r.nombre}"
        fila[1].text = f"{r.porcentaje:.1f}%"
        fila[2].text = BUCKET_ETIQUETAS.get(r.bucket, r.bucket)
        estilos.colorear_texto(
            fila[2].paragraphs[0].runs[0], estilos.COLOR_POR_BUCKET.get(r.bucket, estilos.COLOR_TEXTO_TENUE_HEX)
        )
        if modo_revision:
            colores.sombrear_run(fila[1].paragraphs[0].runs[0], colores.RESULTADO_CALCULADO)
            colores.sombrear_run(fila[2].paragraphs[0].runs[0], colores.RESULTADO_CALCULADO)
    estilos.estilizar_tabla(tabla, alineacion_columnas=["l", "c", "c"])
    estilos.fijar_ancho_columnas(tabla, [85, 45, 30])
    return tabla
